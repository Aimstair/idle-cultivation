import sys
from pycel import ExcelCompiler
from openpyxl import load_workbook
p=sys.argv[1]
xl=ExcelCompiler(filename=p)
wb=load_workbook(p)
errs=[];n=0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith("="):
                n+=1
                try:
                    v=xl.evaluate(f"{ws.title}!{c.coordinate}")
                    if isinstance(v,str) and v.startswith("#"): errs.append((ws.title,c.coordinate,v,c.value))
                except Exception as e:
                    errs.append((ws.title,c.coordinate,repr(e)[:80],c.value))
print("formulas",n,"errors",len(errs))
for e in errs[:20]: print(e)
def show(rng):
    sh,r=rng.split("!")
    print(rng, xl.evaluate(rng))
for ref in ["Checks!C5","Checks!C8","Checks!C9","Checks!C10","Checks!C11","Checks!C12","Checks!C13","Checks!C14","Checks!C15"]:
    show(ref)
for r in range(5,15):
    print([xl.evaluate(f"Checks!{c}{r}") for c in "ABCDEH"])
print("body")
for r in range(20,27):
    print([xl.evaluate(f"Checks!{c}{r}") for c in "ABCDEF"])
print("ext")
for r in range(30,36):
    print([xl.evaluate(f"Checks!{c}{r}") for c in "ABCDE"])
print("other")
for r in range(39,44):
    print([xl.evaluate(f"Checks!{c}{r}") for c in "ABCD"])
print("rebirth")
for r in range(30,36):
    print([round(x,2) if isinstance(x,float) else x for x in [xl.evaluate(f"Rebirth!{c}{r}") for c in "ABCDEFGHIJKLOPQRS"]])
